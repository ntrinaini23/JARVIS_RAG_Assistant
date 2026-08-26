import os
from pypdf import PdfReader
from docx import Document

class DocumentLoader:
    @staticmethod
    def load_txt(file_path: str) -> str:
        """Extracts text from a plain TXT file."""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    @staticmethod
    def load_pdf(file_path: str) -> str:
        """Extracts text from a PDF file page by page."""
        reader = PdfReader(file_path)
        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n\n".join(text_parts)

    @staticmethod
    def load_docx(file_path: str) -> str:
        """Extracts text from a Word DOCX file."""
        doc = Document(file_path)
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text_parts.append(paragraph.text)
        
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text for cell in row.cells if cell.text]
                if row_text:
                    text_parts.append(" | ".join(row_text))
                    
        return "\n".join(text_parts)

    @classmethod
    def load_document(cls, file_path: str) -> str:
        """Determines file type and extracts raw text."""
        _, ext = os.path.splitext(file_path.lower())
        if ext == '.txt':
            return cls.load_txt(file_path)
        elif ext == '.pdf':
            return cls.load_pdf(file_path)
        elif ext == '.docx':
            return cls.load_docx(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")
